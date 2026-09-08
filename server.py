import io
import json
import os
import sqlite3
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional
from defusedxml import ElementTree as ET
from fastapi import FastAPI, File, Form, Header, HTTPException, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, EmailStr
import pandas as pd

import auth

DB_PATH = "banco_notas.db"
PASTA_XMLS_PROCESSADOS = "xmls_processados"

app = FastAPI(
    title="Gestão de Produtos API",
    description="API de alta performance para precificação de produtos e gestão de NF-e",
    version="2.0.0",
)

# CORS middleware para permitir comunicação com o frontend React/Vite
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def inicializar_banco() -> None:
    os.makedirs(PASTA_XMLS_PROCESSADOS, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS historico_nfe (
            chave_acesso TEXT PRIMARY KEY,
            numero TEXT,
            serie TEXT,
            data_emissao TEXT,
            fornecedor TEXT,
            cnpj_emit TEXT,
            valor_total REAL,
            arquivo_original TEXT,
            arquivo_salvo TEXT,
            data_importacao TEXT,
            usuario_id INTEGER
        )
    """)
    cursor.execute("PRAGMA table_info(historico_nfe)")
    colunas = [col[1] for col in cursor.fetchall()]
    if "usuario_id" not in colunas:
        cursor.execute("ALTER TABLE historico_nfe ADD COLUMN usuario_id INTEGER")
    conn.commit()
    conn.close()
    auth.inicializar_tabela_usuarios(DB_PATH)


inicializar_banco()


# ==========================================
# FUNÇÕES DE PROCESSAMENTO DE XML E CÁLCULOS
# ==========================================
def nome_tag(elemento: ET.Element) -> str:
    return elemento.tag.split("}")[-1]


def encontrar_elemento(elemento_pai: Optional[ET.Element], nome: str) -> Optional[ET.Element]:
    if elemento_pai is None:
        return None
    for elemento in elemento_pai.iter():
        if nome_tag(elemento) == nome:
            return elemento
    return None


def encontrar_filho(elemento_pai: Optional[ET.Element], nome: str) -> Optional[ET.Element]:
    if elemento_pai is None:
        return None
    for filho in list(elemento_pai):
        if nome_tag(filho) == nome:
            return filho
    return None


def obter_texto(elemento_pai: Optional[ET.Element], nome: str, padrao: str = "") -> str:
    elemento = encontrar_elemento(elemento_pai, nome)
    if elemento is not None and elemento.text:
        return elemento.text.strip()
    return padrao


def converter_decimal(valor: Any, padrao: Decimal = Decimal("0")) -> Decimal:
    if valor is None:
        return padrao
    texto = str(valor).strip()
    if not texto:
        return padrao
    try:
        return Decimal(texto.replace(",", "."))
    except InvalidOperation:
        return padrao


def valor_tag(elemento_pai: Optional[ET.Element], nome: str) -> Decimal:
    elemento = encontrar_elemento(elemento_pai, nome)
    if elemento is not None and elemento.text:
        return converter_decimal(elemento.text)
    return Decimal("0")


def ler_totais_nfe(root: ET.Element) -> Dict[str, Decimal]:
    total_node = encontrar_elemento(root, "ICMSTot")
    if total_node is None:
        return {
            "frete": Decimal("0"),
            "seguro": Decimal("0"),
            "desconto": Decimal("0"),
            "outras_despesas": Decimal("0"),
        }
    return {
        "frete": valor_tag(total_node, "vFrete"),
        "seguro": valor_tag(total_node, "vSeg"),
        "desconto": valor_tag(total_node, "vDesc"),
        "outras_despesas": valor_tag(total_node, "vOutro"),
    }


def localizar_produtos(root: ET.Element, impostos_selecionados: List[str]) -> List[Dict[str, Any]]:
    produtos = []
    totais_nfe = ler_totais_nfe(root)
    detalhes = []

    for elemento in root.iter():
        if nome_tag(elemento) != "det":
            continue
        produto_node = encontrar_filho(elemento, "prod")
        imposto_node = encontrar_filho(elemento, "imposto")
        if produto_node is None:
            continue
        detalhes.append({"produto_node": produto_node, "imposto_node": imposto_node})

    soma_produtos = sum(valor_tag(item["produto_node"], "vProd") for item in detalhes)
    if soma_produtos <= 0:
        soma_produtos = Decimal("1")

    soma_frete_itens = sum(valor_tag(item["produto_node"], "vFrete") for item in detalhes)
    soma_seguro_itens = sum(valor_tag(item["produto_node"], "vSeg") for item in detalhes)
    soma_desconto_itens = sum(valor_tag(item["produto_node"], "vDesc") for item in detalhes)
    soma_outras_itens = sum(valor_tag(item["produto_node"], "vOutro") for item in detalhes)

    usar_frete_dos_itens = soma_frete_itens > 0
    usar_seguro_dos_itens = soma_seguro_itens > 0
    usar_desconto_dos_itens = soma_desconto_itens > 0
    usar_outras_dos_itens = soma_outras_itens > 0

    for item in detalhes:
        produto_node = item["produto_node"]
        imposto_node = item["imposto_node"]

        codigo = obter_texto(produto_node, "cProd", "Sem código")
        descricao = obter_texto(produto_node, "xProd", "Produto sem descrição")
        unidade = obter_texto(produto_node, "uCom", "UN")

        quantidade = valor_tag(produto_node, "qCom")
        valor_produto = valor_tag(produto_node, "vProd")
        valor_unitario = valor_tag(produto_node, "vUnCom")

        if valor_unitario <= 0 and quantidade > 0:
            valor_unitario = valor_produto / quantidade

        proporcao = valor_produto / soma_produtos

        frete = valor_tag(produto_node, "vFrete") if usar_frete_dos_itens else totais_nfe["frete"] * proporcao
        seguro = valor_tag(produto_node, "vSeg") if usar_seguro_dos_itens else totais_nfe["seguro"] * proporcao
        desconto = valor_tag(produto_node, "vDesc") if usar_desconto_dos_itens else totais_nfe["desconto"] * proporcao
        outras_despesas = valor_tag(produto_node, "vOutro") if usar_outras_dos_itens else totais_nfe["outras_despesas"] * proporcao

        impostos = {
            "ICMS": valor_tag(imposto_node, "vICMS"),
            "ICMS ST": valor_tag(imposto_node, "vICMSST"),
            "FCP": valor_tag(imposto_node, "vFCP"),
            "FCP ST": valor_tag(imposto_node, "vFCPST"),
            "IPI": valor_tag(imposto_node, "vIPI"),
            "II": valor_tag(imposto_node, "vII"),
            "PIS": valor_tag(imposto_node, "vPIS"),
            "COFINS": valor_tag(imposto_node, "vCOFINS"),
        }

        total_impostos_incluidos = sum(
            valor for nome, valor in impostos.items() if nome in impostos_selecionados
        )

        custo_final = (valor_produto + frete + seguro + outras_despesas + total_impostos_incluidos - desconto)
        custo_unitario_final = custo_final / quantidade if quantidade > 0 else Decimal("0")

        registro = {
            "id": f"{codigo}_{datetime.now().timestamp()}",
            "codigo": codigo,
            "produto": descricao,
            "unidade": unidade,
            "quantidade": float(quantidade),
            "valor_produtos": float(valor_produto),
            "custo_unitario_original": float(valor_unitario),
            "frete": float(frete),
            "seguro": float(seguro),
            "desconto": float(desconto),
            "outras_despesas": float(outras_despesas),
            "icms": float(impostos["ICMS"]),
            "icms_st": float(impostos["ICMS ST"]),
            "fcp": float(impostos["FCP"]),
            "fcp_st": float(impostos["FCP ST"]),
            "ipi": float(impostos["IPI"]),
            "ii": float(impostos["II"]),
            "pis": float(impostos["PIS"]),
            "cofins": float(impostos["COFINS"]),
            "impostos_incluidos": float(total_impostos_incluidos),
            "custo_base_final": float(custo_final),
            "custo_unitario_final": float(custo_unitario_final),
            "unidades_por_embalagem": 1.0,
        }
        produtos.append(registro)
    return produtos


def obter_chave_acesso(root: ET.Element) -> Optional[str]:
    inf_nfe = encontrar_elemento(root, "infNFe")
    if inf_nfe is not None:
        id_attr = (inf_nfe.get("Id") or "").strip()
        chave = id_attr.replace("NFe", "").strip()
        if len(chave) == 44 and chave.isdigit():
            return chave
    ch_nfe = obter_texto(root, "chNFe")
    if len(ch_nfe) == 44 and ch_nfe.isdigit():
        return ch_nfe
    return None


def obter_metadados_nfe(root: ET.Element) -> Dict[str, Any]:
    ide_node = encontrar_elemento(root, "ide")
    emit_node = encontrar_elemento(root, "emit")
    total_node = encontrar_elemento(root, "ICMSTot")

    numero = obter_texto(ide_node, "nNF", "")
    serie = obter_texto(ide_node, "serie", "")
    data_emissao = obter_texto(ide_node, "dhEmi", "") or obter_texto(ide_node, "dEmi", "")
    fornecedor = obter_texto(emit_node, "xNome", "Fornecedor não identificado")
    cnpj_emit = obter_texto(emit_node, "CNPJ", "")
    valor_total = float(valor_tag(total_node, "vNF"))

    return {
        "numero": numero,
        "serie": serie,
        "data_emissao": data_emissao[:10] if data_emissao else "",
        "fornecedor": fornecedor,
        "cnpj_emit": cnpj_emit,
        "valor_total": round(valor_total, 2),
    }


def nome_arquivo_padronizado(chave: str, metadados: Dict[str, Any]) -> str:
    numero = "".join(c for c in str(metadados.get("numero", "")) if c.isdigit()) or "SN"
    data = metadados.get("data_emissao", "")
    data_compacta = data.replace("-", "") if data else "sem-data"
    return f"{data_compacta}_NFe{numero}_{chave}.xml"


def salvar_nota_db(dados: Dict[str, Any]) -> None:
    inicializar_banco()
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    cursor.execute("""
        INSERT OR REPLACE INTO historico_nfe 
        (chave_acesso, numero, serie, data_emissao, fornecedor, cnpj_emit, valor_total, arquivo_original, arquivo_salvo, data_importacao, usuario_id)
        VALUES (:chave_acesso, :numero, :serie, :data_emissao, :fornecedor, :cnpj_emit, :valor_total, :arquivo_original, :arquivo_salvo, :data_importacao, :usuario_id)
    """, dados)
    conn.commit()
    conn.close()


def carregar_indice_db(usuario_id: Optional[int] = None) -> List[Dict[str, Any]]:
    inicializar_banco()
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    if usuario_id is not None:
        cursor.execute("""
            SELECT * FROM historico_nfe 
            WHERE usuario_id = ? OR usuario_id IS NULL 
            ORDER BY data_importacao DESC
        """, (usuario_id,))
    else:
        cursor.execute("SELECT * FROM historico_nfe ORDER BY data_importacao DESC")
    linhas = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return linhas


# ==========================================
# MODELOS PYDANTIC PARA AS REQUISIÇÕES
# ==========================================
class LoginRequest(BaseModel):
    identificador: str
    senha: str


class CadastroRequest(BaseModel):
    usuario: str
    nome: str
    email: str
    senha: str
    empresa: Optional[str] = ""
    cargo: Optional[str] = ""


class PerfilUpdateRequest(BaseModel):
    usuario_id: int
    nome: str
    email: str
    empresa: Optional[str] = ""
    cargo: Optional[str] = ""
    markup_padrao: float = 60.0
    custo_adicional_padrao: float = 0.0
    impostos_padrao: List[str] = ["ICMS ST", "FCP ST", "IPI", "II"]


class AlterarSenhaRequest(BaseModel):
    usuario_id: int
    senha_atual: str
    nova_senha: str


class ProdutoItem(BaseModel):
    id: str
    codigo: str
    produto: str
    unidade: str
    quantidade: float
    valor_produtos: float
    custo_unitario_original: float
    frete: float
    seguro: float
    desconto: float
    outras_despesas: float
    icms: float
    icms_st: float
    fcp: float
    fcp_st: float
    ipi: float
    ii: float
    pis: float
    cofins: float
    impostos_incluidos: float
    custo_base_final: float
    custo_unitario_final: float
    unidades_por_embalagem: float = 1.0
    arquivo_xml: Optional[str] = ""


class CalculoRequest(BaseModel):
    produtos: List[Dict[str, Any]]
    markup: float
    custo_adicional_unitario: float
    impostos_selecionados: List[str]
    unidades_por_embalagem: Dict[str, float]


# ==========================================
# ENDPOINTS DE AUTENTICAÇÃO E PERFIL
# ==========================================
@app.post("/api/auth/register")
def register(req: CadastroRequest):
    sucesso, msg = auth.cadastrar_usuario(
        usuario=req.usuario,
        nome=req.nome,
        email=req.email,
        senha=req.senha,
        empresa=req.empresa or "",
        cargo=req.cargo or "",
    )
    if not sucesso:
        raise HTTPException(status_code=400, detail=msg)
    return {"message": msg}


@app.post("/api/auth/login")
def login(req: LoginRequest):
    sucesso, resultado = auth.autenticar_usuario(req.identificador, req.senha)
    if not sucesso:
        raise HTTPException(status_code=401, detail=resultado)
    return {"user": resultado}


@app.get("/api/auth/me/{usuario_id}")
def get_me(usuario_id: int):
    user = auth.obter_usuario_por_id(usuario_id)
    if not user:
        raise HTTPException(status_code=404, detail="Usuário não encontrado.")
    return {"user": user}


@app.put("/api/auth/profile")
def update_profile(req: PerfilUpdateRequest):
    sucesso, msg = auth.atualizar_perfil(
        usuario_id=req.usuario_id,
        nome=req.nome,
        email=req.email,
        empresa=req.empresa or "",
        cargo=req.cargo or "",
        markup_padrao=req.markup_padrao,
        custo_adicional_padrao=req.custo_adicional_padrao,
        impostos_padrao=req.impostos_padrao,
    )
    if not sucesso:
        raise HTTPException(status_code=400, detail=msg)
    user = auth.obter_usuario_por_id(req.usuario_id)
    return {"message": msg, "user": user}


@app.post("/api/auth/change-password")
def change_password(req: AlterarSenhaRequest):
    sucesso, msg = auth.alterar_senha(
        usuario_id=req.usuario_id,
        senha_atual=req.senha_atual,
        nova_senha=req.nova_senha,
    )
    if not sucesso:
        raise HTTPException(status_code=400, detail=msg)
    return {"message": msg}


# ==========================================
# ENDPOINTS DE NF-E E CÁLCULO DE CUSTOS
# ==========================================
@app.post("/api/nfe/upload")
async def upload_xmls(
    files: List[UploadFile] = File(...),
    usuario_id: Optional[int] = Form(None),
    impostos_selecionados: Optional[str] = Form("[\"ICMS ST\", \"FCP ST\", \"IPI\", \"II\"]"),
):
    try:
        impostos_ativos = json.loads(impostos_selecionados) if impostos_selecionados else ["ICMS ST", "FCP ST", "IPI", "II"]
    except Exception:
        impostos_ativos = ["ICMS ST", "FCP ST", "IPI", "II"]

    historico = {n["chave_acesso"]: n for n in carregar_indice_db(usuario_id)}
    todos_produtos = []
    avisos = []
    erros = []
    notas_processadas = []

    for file in files:
        conteudo_bytes = await file.read()
        try:
            root = ET.fromstring(conteudo_bytes)
        except Exception as erro:
            erros.append(f"Erro ao processar {file.filename}: {erro}")
            continue

        chave_acesso = obter_chave_acesso(root)
        metadados = obter_metadados_nfe(root)

        if chave_acesso and chave_acesso in historico:
            avisos.append(f"Nota {metadados.get('numero', '')} ({file.filename}) já havia sido importada anteriormente.")
            continue

        produtos = localizar_produtos(root, impostos_ativos)
        if not produtos:
            erros.append(f"Nenhum produto foi localizado no arquivo {file.filename}.")
            continue

        for prod in produtos:
            prod["arquivo_xml"] = file.filename
            prod["chave_acesso"] = chave_acesso or ""
            prod["numero_nota"] = metadados.get("numero", "")
            prod["fornecedor"] = metadados.get("fornecedor", "")

        todos_produtos.extend(produtos)

        if chave_acesso:
            nome_salvo = nome_arquivo_padronizado(chave_acesso, metadados)
            caminho_salvo = os.path.join(PASTA_XMLS_PROCESSADOS, nome_salvo)
            with open(caminho_salvo, "wb") as f_out:
                f_out.write(conteudo_bytes)

            nova_nota = {
                "chave_acesso": chave_acesso,
                "numero": metadados.get("numero", ""),
                "serie": metadados.get("serie", ""),
                "data_emissao": metadados.get("data_emissao", ""),
                "fornecedor": metadados.get("fornecedor", ""),
                "cnpj_emit": metadados.get("cnpj_emit", ""),
                "valor_total": metadados.get("valor_total", 0),
                "arquivo_original": file.filename,
                "arquivo_salvo": nome_salvo,
                "data_importacao": datetime.now().strftime("%d/%m/%Y %H:%M"),
                "usuario_id": usuario_id,
            }
            salvar_nota_db(nova_nota)
            notas_processadas.append(nova_nota)

    return {
        "produtos": todos_produtos,
        "novas_notas": notas_processadas,
        "avisos": avisos,
        "erros": erros,
    }


@app.post("/api/nfe/calculate")
def recalculate(req: CalculoRequest):
    produtos_calculados = []
    fator_markup = 1 + (req.markup / 100)

    for item in req.produtos:
        item_id = str(item.get("id", ""))
        unidades_embalagem = float(req.unidades_por_embalagem.get(item_id, item.get("unidades_por_embalagem", 1.0)))
        if unidades_embalagem <= 0:
            unidades_embalagem = 1.0

        qtd_original = float(item.get("quantidade", 1.0))
        valor_prod = float(item.get("valor_produtos", 0.0))
        frete = float(item.get("frete", 0.0))
        seguro = float(item.get("seguro", 0.0))
        desconto = float(item.get("desconto", 0.0))
        outras = float(item.get("outras_despesas", 0.0))

        # Impostos detalhados
        impostos_map = {
            "ICMS": float(item.get("icms", 0.0)),
            "ICMS ST": float(item.get("icms_st", 0.0)),
            "FCP": float(item.get("fcp", 0.0)),
            "FCP ST": float(item.get("fcp_st", 0.0)),
            "IPI": float(item.get("ipi", 0.0)),
            "II": float(item.get("ii", 0.0)),
            "PIS": float(item.get("pis", 0.0)),
            "COFINS": float(item.get("cofins", 0.0)),
        }

        impostos_somados = sum(val for nome_imp, val in impostos_map.items() if nome_imp in req.impostos_selecionados)

        # Custo base total
        custo_base = valor_prod + frete + seguro + outras + impostos_somados - desconto
        
        # Ajuste de embalagem e custo adicional
        quantidade_real = qtd_original * unidades_embalagem
        custo_adicional_total = quantidade_real * req.custo_adicional_unitario
        custo_final_total = custo_base + custo_adicional_total

        custo_unitario_final = custo_final_total / quantidade_real if quantidade_real > 0 else 0.0
        preco_revenda_unitario = custo_unitario_final * fator_markup
        total_revenda = quantidade_real * preco_revenda_unitario
        lucro_unitario = preco_revenda_unitario - custo_unitario_final
        lucro_total = quantidade_real * lucro_unitario
        margem_lucro_pct = (lucro_unitario / preco_revenda_unitario * 100) if preco_revenda_unitario > 0 else 0.0

        p_calc = dict(item)
        p_calc.update({
            "unidades_por_embalagem": unidades_embalagem,
            "quantidade_real": round(quantidade_real, 2),
            "impostos_incluidos": round(impostos_somados, 2),
            "custo_adicional_total": round(custo_adicional_total, 2),
            "custo_final": round(custo_final_total, 2),
            "custo_unitario_final": round(custo_unitario_final, 2),
            "preco_revenda_unitario": round(preco_revenda_unitario, 2),
            "total_revenda": round(total_revenda, 2),
            "lucro_unitario": round(lucro_unitario, 2),
            "lucro_total": round(lucro_total, 2),
            "margem_lucro_pct": round(margem_lucro_pct, 2),
        })
        produtos_calculados.append(p_calc)

    # Totais consolidados
    total_custo = sum(p["custo_final"] for p in produtos_calculados)
    total_faturamento = sum(p["total_revenda"] for p in produtos_calculados)
    total_lucro = sum(p["lucro_total"] for p in produtos_calculados)
    total_impostos = sum(p["impostos_incluidos"] for p in produtos_calculados)
    total_frete = sum(p.get("frete", 0.0) for p in produtos_calculados)
    margem_media = (total_lucro / total_faturamento * 100) if total_faturamento > 0 else 0.0

    return {
        "produtos": produtos_calculados,
        "resumo": {
            "total_itens": len(produtos_calculados),
            "total_custo": round(total_custo, 2),
            "total_faturamento": round(total_faturamento, 2),
            "total_lucro": round(total_lucro, 2),
            "total_impostos": round(total_impostos, 2),
            "total_frete": round(total_frete, 2),
            "margem_media_pct": round(margem_media, 2),
        }
    }


@app.get("/api/nfe/history")
def get_history(usuario_id: Optional[int] = None):
    notas = carregar_indice_db(usuario_id)
    return {"notas": notas}


@app.post("/api/nfe/export-csv")
def export_csv(produtos: List[Dict[str, Any]]):
    df = pd.DataFrame(produtos)
    
    # Renomeia colunas para visualização agradável em português no Excel
    mapeamento_colunas = {
        "codigo": "Código",
        "produto": "Produto",
        "unidade": "Unidade Original",
        "quantidade": "Qtd NF-e",
        "unidades_por_embalagem": "Unidades / Embalagem",
        "quantidade_real": "Qtd Real para Venda",
        "valor_produtos": "Valor dos Produtos (R$)",
        "custo_unitario_original": "Custo Unitário Original (R$)",
        "frete": "Frete (R$)",
        "seguro": "Seguro (R$)",
        "desconto": "Desconto (R$)",
        "outras_despesas": "Outras Despesas (R$)",
        "impostos_incluidos": "Total Impostos no Custo (R$)",
        "custo_adicional_total": "Custo Adicional Total (R$)",
        "custo_final": "Custo Final Total (R$)",
        "custo_unitario_final": "Custo Unitário Final (R$)",
        "preco_revenda_unitario": "Preço de Venda Sugerido (R$)",
        "total_revenda": "Faturamento Estimado (R$)",
        "lucro_unitario": "Lucro Unitário (R$)",
        "lucro_total": "Lucro Total Estimado (R$)",
        "margem_lucro_pct": "Margem de Lucro (%)",
        "arquivo_xml": "Arquivo XML",
        "fornecedor": "Fornecedor",
        "numero_nota": "NF-e Nº",
    }

    colunas_presentes = [c for c in mapeamento_colunas.keys() if c in df.columns]
    df_export = df[colunas_presentes].rename(columns=mapeamento_colunas)

    csv_buffer = io.StringIO()
    df_export.to_csv(csv_buffer, index=False, sep=";", decimal=",", encoding="utf-8-sig")
    
    return Response(
        content=csv_buffer.getvalue().encode("utf-8-sig"),
        media_type="text/csv",
        headers={
            "Content-Disposition": f"attachment; filename=precificacao_produtos_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        }
    )


@app.post("/api/nfe/export-excel")
def export_excel(produtos: List[Dict[str, Any]]):
    df = pd.DataFrame(produtos)

    mapeamento_colunas = {
        "codigo": "Código",
        "produto": "Produto",
        "unidade": "Unidade Original",
        "quantidade": "Qtd NF-e",
        "unidades_por_embalagem": "Unidades / Embalagem",
        "quantidade_real": "Qtd Real para Venda",
        "valor_produtos": "Valor dos Produtos (R$)",
        "custo_unitario_original": "Custo Unitário Original (R$)",
        "frete": "Frete (R$)",
        "seguro": "Seguro (R$)",
        "desconto": "Desconto (R$)",
        "outras_despesas": "Outras Despesas (R$)",
        "impostos_incluidos": "Total Impostos no Custo (R$)",
        "custo_adicional_total": "Custo Adicional Total (R$)",
        "custo_final": "Custo Final Total (R$)",
        "custo_unitario_final": "Custo Unitário Final (R$)",
        "preco_revenda_unitario": "Preço de Venda Sugerido (R$)",
        "total_revenda": "Faturamento Estimado (R$)",
        "lucro_unitario": "Lucro Unitário (R$)",
        "lucro_total": "Lucro Total Estimado (R$)",
        "margem_lucro_pct": "Margem de Lucro (%)",
        "arquivo_xml": "Arquivo XML",
        "fornecedor": "Fornecedor",
        "numero_nota": "NF-e Nº",
    }
    colunas_presentes = [c for c in mapeamento_colunas if c in df.columns]
    df_export = df[colunas_presentes].rename(columns=mapeamento_colunas)

    excel_buffer = io.BytesIO()
    with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
        df_export.to_excel(writer, index=False, sheet_name="Precificação")
        worksheet = writer.book["Precificação"]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        header_fill = PatternFill("solid", fgColor="0C8DE4")
        for cell in worksheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for column_cells in worksheet.columns:
            column_letter = column_cells[0].column_letter
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 40)

        for column_name in df_export.columns:
            column_index = df_export.columns.get_loc(column_name) + 1
            if "(R$)" in column_name:
                for cells in worksheet.iter_cols(
                    min_col=column_index, max_col=column_index, min_row=2
                ):
                    for cell in cells:
                        cell.number_format = '"R$" #,##0.00'
            elif "(%)" in column_name:
                for cells in worksheet.iter_cols(
                    min_col=column_index, max_col=column_index, min_row=2
                ):
                    for cell in cells:
                        cell.number_format = '0.00"%"'

    return Response(
        content=excel_buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename=precificacao_produtos_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        },
    )


# Servir build do frontend se existir
if os.path.exists("frontend/dist"):
    app.mount("/", StaticFiles(directory="frontend/dist", html=True), name="frontend")
